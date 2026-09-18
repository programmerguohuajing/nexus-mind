from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from nexusmind.api.server import app
from nexusmind.core.uploads import ingest_uploaded_file

client = TestClient(app)


def test_text_upload_preserves_original_and_creates_markdown(tmp_path):
    result = ingest_uploaded_file(
        "notes.txt",
        "hello knowledge".encode("utf-8"),
        author="Tester",
        vault_root=tmp_path,
    )
    original = tmp_path / result["original_path"]
    markdown = tmp_path / result["markdown_path"]
    assert original.read_text(encoding="utf-8") == "hello knowledge"
    text = markdown.read_text(encoding="utf-8")
    assert "hello knowledge" in text
    assert 'source_type: "text"' in text


def test_duplicate_upload_never_overwrites(tmp_path):
    first = ingest_uploaded_file("same.md", b"# first", vault_root=tmp_path)
    second = ingest_uploaded_file("same.md", b"# second", vault_root=tmp_path)
    assert first["original_path"] != second["original_path"]
    assert first["markdown_path"] != second["markdown_path"]
def test_docx_and_pdf_are_extracted(tmp_path):
    document = Document()
    document.add_heading("DOCX Heading", level=1)
    document.add_paragraph("DOCX body")
    docx_bytes = BytesIO()
    document.save(docx_bytes)

    docx_result = ingest_uploaded_file(
        "sample.docx",
        docx_bytes.getvalue(),
        vault_root=tmp_path,
    )
    docx_md = (tmp_path / docx_result["markdown_path"]).read_text(encoding="utf-8")
    assert "DOCX Heading" in docx_md
    assert "DOCX body" in docx_md
    assert 'source_type: "docx"' in docx_md

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    pdf_bytes = BytesIO()
    writer.write(pdf_bytes)
    pdf_result = ingest_uploaded_file(
        "sample.pdf",
        pdf_bytes.getvalue(),
        vault_root=tmp_path,
    )
    pdf_md = (tmp_path / pdf_result["markdown_path"]).read_text(encoding="utf-8")
    assert "第 1 页" in pdf_md
    assert 'source_type: "pdf"' in pdf_md
def test_upload_api_accepts_multiple_files_without_touching_real_vault(monkeypatch):
    captured = []

    def fake_ingest(filename, data, *, author, source_url):
        captured.append((filename, data, author, source_url))
        return {
            "status": "ingested_successfully",
            "filename": filename,
            "original_path": f"60-References/Files/{filename}",
            "markdown_path": f"60-References/Articles/{filename}.md",
            "version": "abc123",
            "source_type": "text",
            "bytes": len(data),
        }

    monkeypatch.setattr("nexusmind.api.server.ingest_uploaded_file", fake_ingest)
    response = client.post(
        "/api/uploads",
        files=[
            ("files", ("a.txt", b"alpha", "text/plain")),
            ("files", ("b.md", b"# beta", "text/markdown")),
            ("files", ("bad.exe", b"x", "application/octet-stream")),
        ],
        data={"author": "Batch User", "source_url": "https://example.com"},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["total"] == 3
    assert body["succeeded"] == 2
    assert body["failed"] == 1
    assert len(captured) == 2
    assert captured[0][2] == "Batch User"


def test_reference_listing_shows_existing_uploaded_files(tmp_path, monkeypatch):
    files_dir = tmp_path / "60-References" / "Files"
    articles_dir = tmp_path / "60-References" / "Articles"
    meta_dir = tmp_path / "00-Meta"
    files_dir.mkdir(parents=True)
    articles_dir.mkdir(parents=True)
    meta_dir.mkdir(parents=True)

    (files_dir / "sample.pdf").write_bytes(b"%PDF-1.4 sample")
    (articles_dir / "sample.md").write_text(
        "---\n"
        "title: Sample\n"
        "source_file: 60-References/Files/sample.pdf\n"
        "source_type: pdf\n"
        "ingested_at: 2026-09-18 12:00:00\n"
        "---\n# Sample\n",
        encoding="utf-8",
    )
    (meta_dir / "COMPILATION-LOG.md").write_text(
        "`60-References/Articles/sample.md`\n",
        encoding="utf-8",
    )

    monkeypatch.setattr("nexusmind.api.server.VAULT_ROOT", tmp_path)
    response = client.get("/api/references")
    assert response.status_code == 200
    body = response.json()
    assert body["total"] == 1
    item = body["items"][0]
    assert item["filename"] == "sample.pdf"
    assert item["markdown_path"] == "60-References/Articles/sample.md"
    assert item["extracted"] is True
    assert item["compiled"] is True
