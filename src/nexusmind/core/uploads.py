from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple

from docx import Document
from pypdf import PdfReader

from nexusmind.config import VAULT_ROOT
from nexusmind.core.compiler import patch_note

SUPPORTED_EXTENSIONS = {".md", ".txt", ".pdf", ".docx"}


def _safe_name(filename: str) -> str:
    name = Path(filename).name.strip()
    stem = Path(name).stem
    suffix = Path(name).suffix.lower()
    safe_stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", stem).strip(" .")
    return (safe_stem or "upload") + suffix


def _unique_rel_path(folder: str, filename: str, vault_root: Path) -> str:
    base = f"{folder.rstrip('/')}/{filename}"
    if not (vault_root / base).exists():
        return base
    stem = Path(filename).stem
    suffix = Path(filename).suffix
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    candidate = f"{folder.rstrip('/')}/{stem}-{stamp}{suffix}"
    seq = 2
    while (vault_root / candidate).exists():
        candidate = f"{folder.rstrip('/')}/{stem}-{stamp}-{seq}{suffix}"
        seq += 1
    return candidate
def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages: List[str] = []
    for idx, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        pages.append(f"## 第 {idx} 页\n\n{text}" if text else f"## 第 {idx} 页\n\n[未提取到可搜索文本]")
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    blocks: List[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                blocks.append("| " + " | ".join(cells) + " |")
    return "\n\n".join(blocks)


def _extract_text(filename: str, data: bytes) -> Tuple[str, str]:
    ext = Path(filename).suffix.lower()
    if ext == ".md":
        return data.decode("utf-8-sig", errors="replace"), "markdown"
    if ext == ".txt":
        return data.decode("utf-8-sig", errors="replace"), "text"
    if ext == ".pdf":
        return _extract_pdf(data), "pdf"
    if ext == ".docx":
        return _extract_docx(data), "docx"
    raise ValueError(f"Unsupported file type: {ext or '[none]'}")
def ingest_uploaded_file(
    filename: str,
    data: bytes,
    *,
    author: str = "Unknown",
    source_url: str | None = None,
    vault_root: Path = VAULT_ROOT,
) -> Dict[str, Any]:
    """保存原始文件并生成可编译 Markdown，二者都不覆盖已有内容。"""
    safe_name = _safe_name(filename)
    ext = Path(safe_name).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError("Supported file types: .md, .txt, .pdf, .docx")
    if not data:
        raise ValueError("Uploaded file is empty")

    raw_rel = _unique_rel_path("60-References/Files", safe_name, vault_root)
    raw_path = vault_root / raw_rel
    raw_path.parent.mkdir(parents=True, exist_ok=True)
    raw_path.write_bytes(data)

    extracted, kind = _extract_text(safe_name, data)
    title = Path(safe_name).stem
    note_filename = title + ".md"
    note_rel = _unique_rel_path("60-References/Articles", note_filename, vault_root)
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    source_link = raw_rel.replace("\\", "/")
    clean_author = (author or "Unknown").replace('"', "'")
    clean_url = (source_url or "").replace('"', "'")
    markdown = f"""---
title: "{title.replace('"', "'")}"
author: "{clean_author}"
url: "{clean_url}"
ingested_at: "{timestamp}"
source_file: "{source_link}"
source_type: "{kind}"
tags:
  - "#raw-reference"
  - "#uploaded"
---

# {title}

> 原始文件：[[{source_link}]]

{extracted.strip()}
"""
    result = patch_note(
        note_rel,
        markdown,
        if_match="NEW",
        actor="upload",
        privileged=True,
        vault_root=vault_root,
    )
    return {
        "status": "ingested_successfully",
        "filename": safe_name,
        "original_path": raw_rel,
        "markdown_path": note_rel,
        "version": result["version"],
        "source_type": kind,
        "bytes": len(data),
    }


def ingest_uploaded_files(
    files: Iterable[Tuple[str, bytes]],
    *,
    author: str = "Unknown",
    source_url: str | None = None,
    vault_root: Path = VAULT_ROOT,
) -> List[Dict[str, Any]]:
    return [
        ingest_uploaded_file(
            filename,
            data,
            author=author,
            source_url=source_url,
            vault_root=vault_root,
        )
        for filename, data in files
    ]
